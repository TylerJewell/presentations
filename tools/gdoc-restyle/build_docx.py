#!/usr/bin/env python3
"""Build an upload-ready .docx in the AAO white paper style from parse_export.py blocks.

The style tokens come from the AAO white paper print skin
(~/docs/docs/bin/whitepaper/print.css). Text is carried over verbatim; only
structure and formatting change.

    python tools/gdoc-restyle/build_docx.py <blocks.json> <out.docx> --title "Client — Doc Title" [--meta "Akka · September 2026"]

A title containing " — " splits into a teal eyebrow (before) and the gold cover
title (after). The footer carries the full title and page n/N on every page
after the cover. Writes <out.docx> and <out.docx>.b64 for the Drive connector.
"""
import argparse
import base64
import io
import json
import re
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from PIL import Image, ImageDraw

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# print.css :root
INK, MUTED, FOOT = '1a1a1a', '5f6b73', '8a8a8a'
ACCENT, GOLD, GOLD_DK = '02a4a7', 'f5b60b', 'c99a10'
RULE, BORDER_STRONG, INLINE_BG, CARD_BG = 'adadad', 'c9c9c9', 'e5e5e5', 'f6f6f6'
# Docs reads "Instrument Sans Medium" as Instrument Sans at weight 500
SANS, SANS_MEDIUM, MONO = 'Instrument Sans', 'Instrument Sans Medium', 'Roboto Mono'

PAGE_W, MARGIN_LR = Mm(210), Mm(16)
CONTENT_W = PAGE_W - 2 * MARGIN_LR
CONTENT_DXA = int(CONTENT_W / 635)  # 1 dxa = 635 EMU

NUMERIC = re.compile(r'^[~<>$N\d]')
BASE64_WARN = 23000  # uploads of 24.4K base64 characters failed; 23.6K and 22.0K succeeded


def rgb(hex_):
    return RGBColor.from_string(hex_.upper())


def text_of(runs):
    return ''.join(r['t'] for r in runs)


def set_font(target, name, size=None, color=None, bold=None, italic=None, caps=None):
    f = target.font
    f.name = name
    rfonts = target.element.get_or_add_rPr().find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        target.element.get_or_add_rPr().insert(0, rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), name)
    # theme font attributes override the named font
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        rfonts.attrib.pop(qn(a), None)
    if size is not None:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = rgb(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if caps is not None:
        f.all_caps = caps


def para_fmt(pf, before=None, after=None, line=None, keep_next=None):
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next


def add_runs(p, runs, size=None, color=None, force_bold=False):
    for r in runs:
        for i, part in enumerate(r['t'].split('\n')):
            if i:
                p.add_run().add_break()
            if not part:
                continue
            run = p.add_run(part)
            if r['b'] or force_bold:
                run.bold = True
            if r['i']:
                run.italic = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)


def brand_bar():
    """5px gradient #ffce4a -> #04c4c5 -> #d70023 with rounded ends, as PNG bytes."""
    w, h = 2400, 24
    stops = [(0, (0xFF, 0xCE, 0x4A)), (0.5, (0x04, 0xC4, 0xC5)), (1, (0xD7, 0x00, 0x23))]
    grad = Image.new('RGB', (w, h))
    px = grad.load()
    for x in range(w):
        t = x / (w - 1)
        (t0, c0), (t1, c1) = (stops[0], stops[1]) if t <= 0.5 else (stops[1], stops[2])
        u = (t - t0) / (t1 - t0)
        c = tuple(round(c0[k] + (c1[k] - c0[k]) * u) for k in range(3))
        for y in range(h):
            px[x, y] = c
    mask = Image.new('L', (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w - 1, h - 1], radius=14, fill=255)
    img = Image.new('RGBA', (w, h), (0, 0, 0, 0))
    img.paste(grad, (0, 0), mask)
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return buf, h / w


def define_styles(doc):
    st = doc.styles
    set_font(st['Normal'], SANS, 10.5, INK)
    para_fmt(st['Normal'].paragraph_format, before=0, after=7, line=1.55)

    h1 = st['Heading 1']
    set_font(h1, SANS_MEDIUM, 19, INK, bold=False, italic=False)
    para_fmt(h1.paragraph_format, before=26, after=12, line=1.15, keep_next=True)
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '7'), ('w:color', RULE.upper())):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    h1.element.get_or_add_pPr().append(bdr)

    h2 = st['Heading 2']
    set_font(h2, SANS, 13.5, INK, bold=True, italic=False)
    para_fmt(h2.paragraph_format, before=18, after=8, line=1.15, keep_next=True)

    h3 = st['Heading 3']
    set_font(h3, SANS, 11, INK, bold=True, italic=False)
    para_fmt(h3.paragraph_format, before=13, after=6, line=1.15, keep_next=True)

    title = st['Title']
    set_font(title, SANS_MEDIUM, 42, GOLD, bold=False)
    para_fmt(title.paragraph_format, before=0, after=18, line=1.05)
    for el in title.element.pPr.findall(qn('w:pBdr')):
        title.element.pPr.remove(el)

    # the default template colours headings and Title from the theme
    for s in (h1, h2, h3, title):
        for el in s.element.rPr.findall(qn('w:color')):
            el.attrib.pop(qn('w:themeColor'), None)
            el.attrib.pop(qn('w:themeShade'), None)

    for name in ('List Bullet', 'List Number'):
        set_font(st[name], SANS, 10.5, INK)
        para_fmt(st[name].paragraph_format, before=0, after=3.5, line=1.45)

    # Docs ignores paragraph-level tab clears, so the Letter-sized tabs go from the style itself
    st['Footer'].paragraph_format.tab_stops.clear_all()


def tc_pr(cell):
    return cell._tc.get_or_add_tcPr()


def cell_width(cell, dxa):
    w = tc_pr(cell).find(qn('w:tcW'))
    if w is None:
        w = OxmlElement('w:tcW')
        tc_pr(cell).insert(0, w)
    w.set(qn('w:w'), str(dxa))
    w.set(qn('w:type'), 'dxa')


def cell_borders(cell, **sides):
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        if side in sides:
            sz, color = sides[side]
            for k, v in (('w:val', 'single'), ('w:sz', str(sz)), ('w:space', '0'), ('w:color', color.upper())):
                el.set(qn(k), v)
        else:
            el.set(qn('w:val'), 'nil')
        b.append(el)
    tc_pr(cell).append(b)


def cell_shade(cell, fill):
    shd = OxmlElement('w:shd')
    for k, v in (('w:val', 'clear'), ('w:color', 'auto'), ('w:fill', fill.upper())):
        shd.set(qn(k), v)
    tc_pr(cell).append(shd)


def cell_margins(cell, top, right, bottom, left):
    mar = OxmlElement('w:tcMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(int(v * 20)))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tc_pr(cell).append(mar)


def table_setup(table, col_dxa):
    tblpr = table._tbl.tblPr
    for tag, attrs in (('w:tblW', {'w:w': str(sum(col_dxa)), 'w:type': 'dxa'}),
                       ('w:tblLayout', {'w:type': 'fixed'})):
        el = tblpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tblpr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 5), ('left', 8), ('bottom', 5), ('right', 8)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(v * 20))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblpr.append(mar)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tblpr.append(borders)
    for gc, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), col_dxa):
        gc.set(qn('w:w'), str(w))


def numeric_cols(rows, ncols):
    """Columns after the row-label column where 70% of cells are short figures."""
    out = set()
    for c in range(1, ncols):
        vals = [text_of(r['cells'][c]['paras'][0]).strip() for r in rows
                if r['kind'] != 'head' and c < len(r['cells'])]
        vals = [v for v in vals if v]
        hits = [v for v in vals if len(v) <= 16 and NUMERIC.match(v) and re.search(r'\d', v)]
        if vals and len(hits) / len(vals) >= 0.7:
            out.add(c)
    return out


def column_widths(rows, ncols):
    """Width follows content: the longest cell (capped at 40 chars) or the longest word, whichever is larger."""
    widths = []
    for c in range(ncols):
        body = [text_of(r['cells'][c]['paras'][0]) for r in rows if r['kind'] != 'head' and c < len(r['cells'])]
        words = [w for t in body + [text_of(rows[0]['cells'][c]['paras'][0])] for w in t.split()]
        widths.append(max(max((len(w) for w in words), default=4) * 1.45 + 4,
                          min(max((len(t) for t in body), default=4), 40)))
    col = [int(w * CONTENT_DXA / sum(widths)) for w in widths]
    col[-1] += CONTENT_DXA - sum(col)
    return col


def spacer(doc, after):
    para_fmt(doc.add_paragraph().paragraph_format, before=0, after=after, line=1.0)


def data_table(doc, rows):
    ncols = max(len(r['cells']) for r in rows)
    col_dxa = column_widths(rows, ncols)
    nums = numeric_cols(rows, ncols)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_setup(table, col_dxa)
    for ri, r in enumerate(rows):
        kind = r['kind']
        # a row that is bold end to end is a total row even without the grey fill
        if kind == 'body' and all(c['paras'][0] and all(x['b'] for x in c['paras'][0]) for c in r['cells']):
            kind = 'total'
        tr_pr = table.rows[ri]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if kind == 'head':
            tr_pr.append(OxmlElement('w:tblHeader'))
        for ci in range(ncols):
            cell = table.rows[ri].cells[ci]
            cell_width(cell, col_dxa[ci])
            v = OxmlElement('w:vAlign')
            v.set(qn('w:val'), 'top')
            tc_pr(cell).append(v)
            if kind == 'head':
                cell_shade(cell, CARD_BG)
                cell_borders(cell, bottom=(12, GOLD_DK))
            elif kind == 'total':
                cell_borders(cell, top=(8, INK), bottom=(6, BORDER_STRONG))
            else:
                cell_borders(cell, bottom=(6, BORDER_STRONG))
            src = r['cells'][ci] if ci < len(r['cells']) else {'paras': [[]]}
            for pi, runs in enumerate(src['paras']):
                p = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
                # Docs ignores keep-with-next inside tables; Word honours it
                para_fmt(p.paragraph_format, before=0, after=0, line=1.25, keep_next=ri < len(rows) - 1)
                if ci in nums:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_runs(p, runs, size=9.5, force_bold=kind in ('head', 'total'))
    spacer(doc, 4)


def callout(doc, label_runs, items, bar_color):
    """Single-cell table: grey fill, coloured bar on the left. items are ('bullet'|'para', runs)."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table_setup(table, [CONTENT_DXA])
    cell = table.rows[0].cells[0]
    cell_width(cell, CONTENT_DXA)
    cell_shade(cell, INLINE_BG)
    cell_borders(cell, left=(18, bar_color))
    cell_margins(cell, 11, 14, 11, 14)
    paras = []
    if label_runs:
        paras.append((None, label_runs, 4))
    paras += [('List Bullet' if kind == 'bullet' else None, runs, 3 if kind == 'bullet' else 0) for kind, runs in items]
    for pi, (style, runs, after) in enumerate(paras):
        p = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
        if style:
            p.style = doc.styles[style]
        para_fmt(p.paragraph_format, before=0, after=after, line=1.45)
        add_runs(p, runs)
    spacer(doc, 6)


def hyphen_items(blocks, i):
    """Paragraphs typed as '- item' after index i, as bullet items, and the index of the last one."""
    items = []
    while i + 1 < len(blocks) and blocks[i + 1]['k'] == 'p' and text_of(blocks[i + 1]['runs']).startswith('- '):
        i += 1
        runs = [dict(x) for x in blocks[i]['runs']]
        runs[0]['t'] = runs[0]['t'][2:]
        items.append(('bullet', runs))
    return items, i


def body(doc, blocks):
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b['k'] == 'h':
            doc.add_paragraph(b['text'].strip(), style=f"Heading {b['lvl']}")
        elif b['k'] == 'p':
            t = text_of(b['runs']).strip()
            if t.startswith('[Insert > Table of contents'):
                pass  # author's placeholder; Docs inserts a native TOC
            elif b.get('hl') and t.endswith(':'):
                items, i = hyphen_items(blocks, i)
                callout(doc, b['runs'], items, GOLD_DK)
            elif t.startswith('Note:'):
                callout(doc, None, [('para', b['runs'])], ACCENT)
            elif t.startswith('- '):
                items, i = hyphen_items(blocks, i - 1)
                for _, runs in items:
                    add_runs(doc.add_paragraph(style='List Bullet'), runs)
            elif all(x['i'] for x in b['runs']):
                p = doc.add_paragraph()
                para_fmt(p.paragraph_format, before=0, after=6, line=1.4)
                add_runs(p, b['runs'], size=9, color=MUTED)
            else:
                add_runs(doc.add_paragraph(), b['runs'])
        elif b['k'] == 'list':
            style = 'List Number' if b['ordered'] else 'List Bullet'
            for it in b['items']:
                add_runs(doc.add_paragraph(style=style), it['runs'])
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
        elif b['k'] == 'table':
            data_table(doc, b['rows'])
        i += 1


def cover(doc, title, meta):
    bar, ratio = brand_bar()
    p = doc.add_paragraph()
    para_fmt(p.paragraph_format, before=24, after=52, line=1.0)
    p.add_run().add_picture(bar, width=CONTENT_W, height=Emu(int(CONTENT_W * ratio)))
    eyebrow, _, main = title.rpartition(' — ')
    if eyebrow:
        p = doc.add_paragraph()
        para_fmt(p.paragraph_format, before=0, after=16, line=1.2)
        set_font(p.add_run(eyebrow), MONO, 11, ACCENT, caps=True)
    p = doc.add_paragraph(main, style='Title')
    if meta:
        p = doc.add_paragraph()
        para_fmt(p.paragraph_format, before=44, after=0, line=1.2)
        set_font(p.add_run(meta), MONO, 9.5, MUTED, caps=True)
    p.add_run().add_break(WD_BREAK.PAGE)


def footer(section, title):
    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(CONTENT_W - Pt(1), WD_TAB_ALIGNMENT.RIGHT)
    para_fmt(fp.paragraph_format, before=0, after=0, line=1.0)
    for text in (title, '\t'):
        set_font(fp.add_run(text), MONO, 7, FOOT)
    for instr, after in (('PAGE', '/'), ('NUMPAGES', None)):
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), instr)
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), MONO)
        col = OxmlElement('w:color')
        col.set(qn('w:val'), FOOT.upper())
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '14')
        rpr.extend([rf, col, sz])
        t = OxmlElement('w:t')
        t.text = '1'
        r.extend([rpr, t])
        fld.append(r)
        fp._p.append(fld)
        if after:
            set_font(fp.add_run(after), MONO, 7, FOOT)


# Parts Google Docs does not need. Dropping them keeps the base64 upload short.
DROP = {'word/stylesWithEffects.xml', 'docProps/thumbnail.jpeg', 'docProps/app.xml',
        'word/theme/theme1.xml', 'word/webSettings.xml', 'word/settings.xml', 'word/fontTable.xml',
        'customXml/item1.xml', 'customXml/_rels/item1.xml.rels', 'customXml/itemProps1.xml'}
DROP_RELS = re.compile(r'<Relationship [^>]*Target="[^"]*(stylesWithEffects|thumbnail|customXml|'
                       r'theme1|webSettings|settings\.xml|fontTable|app\.xml)[^"]*"[^>]*/>')


def slim(docx_bytes):
    """Remove unused parts and styles from the default python-docx template."""
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes))
    xml = ''.join(zin.read(n).decode('utf-8') for n in zin.namelist()
                  if n == 'word/document.xml' or n.startswith('word/footer'))
    styles = zin.read('word/styles.xml').decode('utf-8')
    by_id = {re.search(r'w:styleId="([^"]+)"', s).group(1): s
             for s in re.findall(r'<w:style [^>]*>.*?</w:style>', styles, re.S)}
    keep = set(re.findall(r'w:(?:pStyle|rStyle|tblStyle) w:val="([^"]+)"', xml))
    keep |= {'Normal', 'DefaultParagraphFont', 'TableNormal', 'NoList', 'Footer', 'Header',
             'Heading1', 'Heading2', 'Heading3', 'Title', 'ListBullet', 'ListNumber'}
    grew = True
    while grew:  # follow basedOn / link / next
        grew = False
        for sid in list(keep):
            for ref in re.findall(r'<w:(?:basedOn|link|next) w:val="([^"]+)"', by_id.get(sid, '')):
                if ref in by_id and ref not in keep:
                    keep.add(ref)
                    grew = True
    for sid, s in by_id.items():
        if sid not in keep:
            styles = styles.replace(s, '')
    styles = re.sub(r'<w:latentStyles.*?</w:latentStyles>', '', styles, flags=re.S)

    out = io.BytesIO()
    zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED, compresslevel=9)
    for info in zin.infolist():
        n = info.filename
        if n in DROP:
            continue
        data = zin.read(n)
        if n == 'word/styles.xml':
            data = styles.encode('utf-8')
        elif n == '[Content_Types].xml':
            s = data.decode('utf-8')
            for d in DROP:
                s = re.sub(r'<Override PartName="/' + re.escape(d) + r'"[^>]*/>', '', s)
            data = s.encode('utf-8')
        elif n.endswith('.rels'):
            data = DROP_RELS.sub('', data.decode('utf-8')).encode('utf-8')
        zout.writestr(n, data)
    zout.close()
    return out.getvalue()


def build(blocks, title, meta):
    doc = Document()
    doc.core_properties.title = title
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = MARGIN_LR
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(18)
    sec.footer_distance = Mm(8)
    sec.different_first_page_header_footer = True
    define_styles(doc)
    cover(doc, title, meta)
    body(doc, blocks)
    footer(sec, title)
    buf = io.BytesIO()
    doc.save(buf)
    return slim(buf.getvalue())


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('blocks')
    ap.add_argument('out')
    ap.add_argument('--title', required=True, help='Drive title of the source doc')
    ap.add_argument('--meta', help='cover line under the title, e.g. "Akka · September 2026"')
    args = ap.parse_args()
    data = build(json.load(open(args.blocks, encoding='utf-8')), args.title, args.meta)
    open(args.out, 'wb').write(data)
    b64 = base64.b64encode(data).decode('ascii')
    open(args.out + '.b64', 'w').write(b64)
    print(f'{args.out}: {len(data)} bytes, {len(b64)} base64 chars')
    if len(b64) > BASE64_WARN:
        print(f'warning: base64 over {BASE64_WARN} chars; the connector upload may be rejected')
